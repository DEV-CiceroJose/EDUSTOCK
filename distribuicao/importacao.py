import csv
import io
import unicodedata
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
from defusedxml.common import DefusedXmlException

from django.core.exceptions import ValidationError
from openpyxl import load_workbook
from core.models import Turma
from .models import Aluno

MAX_LINHAS = 10000


def normalizar(valor):
    return ' '.join(unicodedata.normalize('NFKC', str(valor or '')).strip().casefold().split())


def ler_arquivo(arquivo):
    if not arquivo or arquivo.size > 5 * 1024 * 1024:
        raise ValidationError('Selecione um arquivo CSV ou XLSX de até 5 MB.')
    try:
        if arquivo.name.lower().endswith('.xlsx'):
            with ZipFile(arquivo) as zip_file:
                if sum(i.file_size for i in zip_file.infolist()) > 30 * 1024 * 1024:
                    raise ValidationError('Planilha descompactada muito grande.')
            arquivo.seek(0)
            livro = load_workbook(arquivo, read_only=True, data_only=False, keep_links=False)
            try:
                planilha = livro.active
                if planilha is None:
                    raise ValidationError('A planilha precisa conter uma aba de alunos.')
                if (planilha.max_column or 0) > 20 or (planilha.max_row or 0) > MAX_LINHAS + 1:
                    raise ValidationError('Limite: 10.000 alunos e 20 colunas.')
                # As dimensões declaradas podem estar ausentes ou incorretas.
                planilha.reset_dimensions()
                linhas = []
                for numero, linha in enumerate(planilha.iter_rows(values_only=True), 1):
                    if numero > MAX_LINHAS + 1 or len(linha) > 20:
                        raise ValidationError('Limite: 10.000 alunos e 20 colunas.')
                    linhas.append(linha)
            finally:
                livro.close()
        elif arquivo.name.lower().endswith('.csv'):
            texto = arquivo.read().decode('utf-8-sig')
            delimitador = ';' if ';' in texto.splitlines()[0] else ','
            linhas = list(csv.reader(io.StringIO(texto), delimiter=delimitador))
        else:
            raise ValidationError('Use o modelo XLSX ou CSV em UTF-8.')
    except (ValueError, UnicodeError, BadZipFile, KeyError, IndexError, OSError, ParseError, DefusedXmlException) as exc:
        raise ValidationError('Arquivo inválido. Baixe e preencha o modelo novamente.') from exc
    if not linhas or len(linhas) > MAX_LINHAS + 1:
        raise ValidationError('Planilha vazia ou acima de 10.000 alunos.')
    cabecalho = [normalizar(c).replace('é', 'e') for c in linhas[0]]
    if any(cabecalho.count(c) != 1 for c in ['nome', 'turma', 'serie', 'sexo']):
        raise ValidationError('As colunas obrigatórias são nome, turma, série e sexo, sem repetição.')
    return [dict(zip(cabecalho, linha)) for linha in linhas[1:] if any(v is not None and str(v).strip() for v in linha)]


def preparar_linhas(linhas, escola):
    turmas = {}
    for turma in Turma.objects.filter(escola=escola, ativo=True):
        turmas.setdefault(normalizar(turma.nome), []).append(turma)
    existentes = {}
    for aluno in Aluno.objects.filter(escola=escola):
        existentes.setdefault(normalizar(aluno.nome), []).append(aluno)
    vistos, novos, erros, avisos = set(), [], [], []
    sexos = {'f': 'F', 'feminino': 'F', 'm': 'M', 'masculino': 'M', 'n': 'N', 'não informado': 'N', 'nao informado': 'N'}
    for numero, linha in enumerate(linhas, 2):
        nome, serie = str(linha.get('nome') or '').strip(), str(linha.get('serie') or '').strip()
        turma_match = turmas.get(normalizar(linha.get('turma')), [])
        sexo = sexos.get(normalizar(linha.get('sexo')))
        if not nome or len(nome) > 200 or not serie or len(serie) > 40 or not sexo or len(turma_match) != 1 or nome.startswith(('=', '+', '-', '@')):
            erros.append(f'Linha {numero}: confira nome, série, sexo (F/M/N) e turma existente na escola.')
            continue
        turma = turma_match[0]
        chave = normalizar(nome)
        if chave in vistos:
            erros.append(f'Linha {numero}: nome repetido no arquivo; confira os homônimos e cadastre-os individualmente.')
            continue
        vistos.add(chave)
        anteriores = existentes.get(chave, [])
        if anteriores:
            if len(anteriores) == 1 and anteriores[0].turma_id == turma.pk and anteriores[0].serie == serie and anteriores[0].sexo == sexo and anteriores[0].ativo:
                avisos.append(f'Linha {numero}: {nome} já está cadastrado; será ignorado. Se for um homônimo, cadastre individualmente.')
            else:
                erros.append(f'Linha {numero}: nome já cadastrado com dados diferentes ou ambíguos; revise pelo ID no cadastro.')
            continue
        novos.append({'nome': nome, 'serie': serie, 'sexo': sexo, 'turma_id': turma.pk})
    return novos, erros, avisos
