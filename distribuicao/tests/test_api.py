import io
from datetime import timedelta

from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from openpyxl import Workbook
from rest_framework.test import APIClient

from plataforma.models import Perfil, TokenAcesso, VinculoUsuario, Escola
from core.models import Turma, PinAcesso, LoteEstoque
from core.operacao_auth import criar_token
from distribuicao.models import Aluno, Destinatario, Rodada, ItemRodada
from .test_entregas import EntregasTest


class ApiTest(EntregasTest):
    def setUp(self):
        super().setUp()
        cache.clear()
        Perfil.objects.create(user=self.user, papel='ADMIN')
        self.vinculo = VinculoUsuario.objects.create(user=self.user, escola=self.escola,
                           municipio=self.escola.municipio, papel='GESTOR_ESCOLA')
        _, token = TokenAcesso.emitir(user=self.user, escola=self.escola, municipio=self.escola.municipio,
                                      papel_rede='GESTOR_ESCOLA', expira_em=timezone.now() + timedelta(hours=1))
        self.client = APIClient()
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {token}')
        self.op = APIClient()
        self.op.credentials(HTTP_X_OPERACAO_TOKEN=criar_token('ALUNO_REP', turma=self.turma.nome,
            turma_id=self.turma.pk, pin_acesso_id=self.pin.pk, pin_versao=self.pin.pin_fingerprint, escola_id=self.escola.pk))

    def previa(self, texto):
        arquivo = SimpleUploadedFile('alunos.csv', texto.encode('utf-8'))
        return self.client.post('/api/distribuicao/importacao/', {'arquivo': arquivo}, format='multipart')

    def test_importacao_preview_sem_escrever_e_confirmacao_idempotente(self):
        r = self.previa('nome;turma;série;sexo\nMaria;1 A;1;F\nAna;1 A;1;F')
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data['novos'], 1)
        self.assertEqual(r.data['ignorados'], 1)
        self.assertEqual(Aluno.objects.count(), 2)
        with self.captureOnCommitCallbacks(execute=True):
            confirmacao = self.client.put('/api/distribuicao/importacao/', {'token': r.data['token']}, format='json')
        self.assertEqual(confirmacao.status_code, 200, confirmacao.data)
        self.assertEqual(Aluno.objects.count(), 3)
        self.assertEqual(self.client.put('/api/distribuicao/importacao/', {'token': r.data['token']}, format='json').status_code, 400)

    def test_linha_invalida_bloqueia_todo_lote(self):
        r = self.previa('nome;turma;serie;sexo\nMaria;1 A;1;F\nPedro;Outra;1;M')
        self.assertIsNone(r.data['token'])
        self.assertTrue(r.data['erros'])
        self.assertEqual(Aluno.objects.count(), 2)

    def test_homonimos_e_dados_divergentes_exigem_revisao(self):
        r = self.previa('nome;turma;serie;sexo\nMaria;1 A;1;F\nMaria;1 A;1;F\nAna;1 A;2;F')
        self.assertEqual(len(r.data['erros']), 2)

    def test_excel_real_e_modelo(self):
        livro = Workbook()
        livro.active.append(['nome', 'turma', 'série', 'sexo'])
        livro.active.append(['Maria', '1 A', '1', 'Feminino'])
        arquivo = io.BytesIO()
        livro.save(arquivo)
        r = self.client.post('/api/distribuicao/importacao/', {'arquivo': SimpleUploadedFile('teste.xlsx', arquivo.getvalue())}, format='multipart')
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data['novos'], 1)
        self.assertEqual(self.client.get('/api/distribuicao/modelo/').status_code, 200)

    def test_excel_sem_dimensoes_de_streaming(self):
        livro = Workbook(write_only=True)
        planilha = livro.create_sheet()
        planilha.append(['nome', 'turma', 'série', 'sexo'])
        planilha.append(['Maria', '1 A', '1', 'F'])
        arquivo = io.BytesIO()
        livro.save(arquivo)
        resposta = self.client.post('/api/distribuicao/importacao/', {
            'arquivo': SimpleUploadedFile('stream.xlsx', arquivo.getvalue())
        }, format='multipart')
        self.assertEqual(resposta.status_code, 200, resposta.data)
        self.assertEqual(resposta.data['novos'], 1)

    def test_pin_transferido_nao_le_lista_anterior(self):
        self.liberar()
        outra = Turma.objects.create(escola=self.escola, nome='2 A', curso='DS', ano=2)
        PinAcesso.objects.filter(pk=self.pin.pk).update(turma=outra)
        self.assertEqual(self.op.get('/api/distribuicao/operacao/').status_code, 401)

    def test_pin_convertido_para_cozinha_nao_le_alunos(self):
        self.liberar()
        PinAcesso.objects.filter(pk=self.pin.pk).update(papel=PinAcesso.COZINHA, turma=None)
        self.assertEqual(self.op.get('/api/distribuicao/operacao/').status_code, 401)

    def test_operador_nao_pode_gerir_mesmo_com_perfil_admin(self):
        self.vinculo.papel = 'OPERADOR'
        self.vinculo.save()
        self.assertEqual(self.client.get('/api/distribuicao/alunos/').status_code, 403)

    def test_api_nao_aceita_turma_ou_produto_de_outra_escola(self):
        outra = Escola.objects.create(nome='Outra', slug='outra', municipio=self.escola.municipio)
        turma = Turma.objects.create(escola=outra, nome='2 A', curso='DS', ano=2)
        r = self.client.post('/api/distribuicao/alunos/', {'nome': 'Maria', 'turma': turma.pk, 'serie': '1', 'sexo': 'F'}, format='json')
        self.assertEqual(r.status_code, 400)
        self.produto.escola = outra
        self.produto.save()
        r = self.client.post('/api/distribuicao/rodadas/', {'titulo': 'Teste', 'tipo': 'KIT_ESCOLAR', 'itens': [{'produto': self.produto.pk, 'quantidade': '1'}]}, format='json')
        self.assertEqual(r.status_code, 400)

    def test_cria_rodada_por_api_valida_quantidade_e_libera(self):
        body = {'titulo': 'Kits', 'tipo': 'KIT_ESCOLAR', 'itens': [{'produto': self.produto.pk, 'quantidade': '1'}]}
        r = self.client.post('/api/distribuicao/rodadas/', body, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        liberada = self.client.post(f"/api/distribuicao/rodadas/{r.data['id']}/estado/", {'estado': 'LIBERADA'}, format='json')
        self.assertEqual(liberada.status_code, 200, liberada.data)
        self.assertEqual(liberada.data['total'], 2)
        body['itens'][0]['quantidade'] = '0'
        self.assertEqual(self.client.post('/api/distribuicao/rodadas/', body, format='json').status_code, 400)

    def test_pin_entrega_relatorio_filtra_e_csv_protege_formulas(self):
        self.menina.nome = '=HYPERLINK("x")'
        self.menina.save()
        self.liberar()
        r = self.op.get('/api/distribuicao/operacao/')
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(len(r.data['alunos']), 1)
        self.assertNotIn('sexo', r.data['alunos'][0])
        alvo = r.data['alunos'][0]['id']
        self.assertEqual(self.op.post('/api/distribuicao/operacao/', {'destinatario': alvo}, format='json').status_code, 200)
        self.assertEqual(self.op.post('/api/distribuicao/operacao/', {'destinatario': alvo}, format='json').status_code, 200)
        r = self.client.get('/api/distribuicao/relatorio/?situacao=PENDENTE')
        self.assertEqual(r.data['entregues'], 1)
        self.assertEqual(r.data['resultados'], [])
        r = self.client.get('/api/distribuicao/relatorio/?exportar=csv')
        self.assertIn("'=HYPERLINK", r.content.decode())
        self.assertEqual(r['Cache-Control'], 'no-store')

    def test_pin_revogado_cozinha_e_outros_nao_acessam(self):
        self.liberar()
        self.pin.ativo = False
        self.pin.save()
        self.assertEqual(self.op.get('/api/distribuicao/operacao/').status_code, 401)
        cozinha = APIClient()
        cozinha.credentials(HTTP_X_OPERACAO_TOKEN=criar_token('COZINHA', escola_id=self.escola.pk))
        self.assertEqual(cozinha.get('/api/distribuicao/operacao/').status_code, 403)
        self.assertEqual(self.client.get('/api/distribuicao/operacao/').status_code, 403)
        self.assertEqual(APIClient().get('/api/distribuicao/relatorio/').status_code, 401)

    def test_outra_turma_nao_ve_nomes_e_nem_contagens(self):
        turma = Turma.objects.create(escola=self.escola, nome='2 A', curso='DS', ano=2)
        Aluno.objects.create(escola=self.escola, turma=turma, nome='Segredo', serie='2', sexo='F')
        self.liberar()
        r = self.op.get('/api/distribuicao/operacao/')
        self.assertEqual(len(r.data['alunos']), 1)
        self.assertNotIn('total', r.data['rodadas'][0])
        outro = Destinatario.objects.get(nome='Segredo')
        self.assertEqual(self.op.post('/api/distribuicao/operacao/', {'destinatario': outro.pk}, format='json').status_code, 400)

    def test_rollback_multiplos_itens(self):
        produto = type(self.produto).objects.create(escola=self.escola, grupo=self.produto.grupo, nome='Caderno', unidade='UN', quantidade=0)
        ItemRodada.objects.create(rodada=self.rodada, produto=produto, quantidade=1)
        self.liberar()
        r = self.op.post('/api/distribuicao/operacao/', {'destinatario': Destinatario.objects.get().pk}, format='json')
        self.assertEqual(r.status_code, 400)
        self.produto.refresh_from_db()
        self.assertEqual(self.produto.quantidade, 10)
        self.assertEqual(LoteEstoque.objects.get(produto=self.produto).quantidade, 10)
        self.assertEqual(Destinatario.objects.get().movimentos.count(), 0)
